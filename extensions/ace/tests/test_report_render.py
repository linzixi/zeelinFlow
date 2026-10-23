from __future__ import annotations

import base64
from io import BytesIO

from docx import Document

from ace_ext.reporting.models import Narrative
from ace_ext.reporting.render import chart, render


def test_report_exports_exact_counts_escaped_sources_and_editable_tables():
    narrative = Narrative(
        title="项目验证报告",
        executive_summary=[{"text": "负面需复核", "query_ids": ["q1"]}],
        recommendations=[{"text": "核对事实", "query_ids": ["q1"]}],
        limitations=["短窗口"],
    )
    dataset = {
        "id": "dataset",
        "window": {"start": "2026-10-22T23:30:00+08:00", "end": "2026-10-22T23:35:00+08:00"},
    }
    audits = [
        {
            "id": "q1",
            "spec": {"operation": "overview"},
            "result": {
                "total": 3,
                "sentiments": [{"label": "negative", "count": 1}, {"label": "neutral", "count": 2}],
                "limitations": ["无传播关系"],
            },
        },
        {
            "id": "q2",
            "spec": {"operation": "timeseries"},
            "result": {
                "points": [
                    {"time": "2026-10-22T23:30:00+08:00", "sentiment": "negative", "count": 1}
                ]
            },
        },
    ]
    artifacts = render(narrative, dataset, audits)
    md = next(a["content"] for a in artifacts if a["name"].endswith(".md"))
    assert "33.33%" in md and "66.67%" in md and "无传播关系" in md
    svg = chart([{"label": "<script>alert(1)</script>", "count": 2}], "来源")
    assert "<script>" not in svg and "&lt;script&gt;" in svg
    word = next(a for a in artifacts if a["name"].endswith(".docx"))
    document = Document(BytesIO(base64.b64decode(word["content_base64"])))
    assert len(document.tables) == 2
    assert "33.33%" == document.tables[0].rows[1].cells[2].text
