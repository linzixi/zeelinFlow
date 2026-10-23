from __future__ import annotations

import base64
from io import BytesIO

from docx import Document

from ace_ext.core.contracts import Envelope
from ace_ext.execution.result_contract import validate_result


def test_report_artifact_contains_editable_word_text_and_source():
    output = validate_result(
        {"operation": "report", "config": {}},
        {"records": []},
        Envelope(
            artifacts=[
                {
                    "name": "日报.md",
                    "content": "# 舆情日报\n\n## 已核验事件\n- 模拟案例\n来源：https://example.invalid/evidence",
                }
            ]
        ),
    )
    word = next(item for item in output.artifacts if item["name"].endswith(".docx"))
    document = Document(BytesIO(base64.b64decode(word["content_base64"])))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "舆情日报" in text and "模拟案例" in text and "example.invalid/evidence" in text
    assert document.paragraphs[0].style.name == "Heading 1"
