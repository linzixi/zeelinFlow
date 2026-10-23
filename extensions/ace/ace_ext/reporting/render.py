from __future__ import annotations

import base64
import html
import json
from collections import defaultdict
from io import BytesIO
from urllib.parse import urlparse

from docx import Document
from docx.shared import Pt

LABELS = {"negative": "负面", "positive": "正面", "neutral": "中性", "unknown": "未判断"}
COLORS = {"negative": "#db6c5b", "positive": "#43a38b", "neutral": "#5b7cfa", "unknown": "#9aa4b5"}


def rows_table(headers, rows):
    def clean(value):
        return str(value).replace("|", "／").replace("\n", " ")

    return "\n".join(
        [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
            *["| " + " | ".join(clean(value) for value in row) + " |" for row in rows],
        ]
    )


def chart(groups, title):
    groups = groups[:12]
    width, height = 800, 92 + len(groups) * 42
    maximum = max((g["count"] for g in groups), default=1) or 1
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" '
        f'role="img" aria-label="{html.escape(title, quote=True)}">',
        '<rect width="100%" height="100%" fill="#f5f7fc" rx="10"/>',
        f'<text x="24" y="32" font-size="19" font-weight="600">{html.escape(title)}</text>',
    ]
    for i, group in enumerate(groups):
        label, count = group["label"], group["count"]
        y, bar = 60 + i * 42, 490 * count / maximum
        parts += [
            f'<text x="24" y="{y + 19}" font-size="13">'
            f"{html.escape(LABELS.get(label, label)[:18])}</text>",
            f'<rect x="210" y="{y}" width="{bar:.2f}" height="26" rx="4" '
            f'fill="{COLORS.get(label, "#5b7cfa")}"/>',
            f'<text x="{220 + bar:.2f}" y="{y + 19}" font-size="13">{count}</text>',
        ]
    return "".join(parts) + "</svg>"


def query(audits, operation, dimension=None):
    return next(
        (
            q
            for q in audits
            if q["spec"]["operation"] == operation
            and (dimension is None or q["spec"].get("dimension") == dimension)
            and not any(q["spec"].get(k) for k in ("sentiment", "keyword", "article_ids", "offset"))
        ),
        None,
    )


def sections(narrative, dataset, audits):
    overview = query(audits, "overview")["result"]
    total = overview["total"]
    sentiments = [
        [
            LABELS.get(g["label"], g["label"]),
            g["count"],
            f"{100 * g['count'] / total:.2f}%" if total else "—",
        ]
        for g in overview["sentiments"]
    ]
    result = [
        (
            "数据范围与口径",
            [
                f"入库时间：{dataset['window']['start']} 至 "
                f"{dataset['window']['end']}（左闭右开）。",
                f"项目去重文章 {total} 篇，数据集 {dataset['id']}。"
                "统计冻结至文章版本，人工复核结论优先。",
                "数据范围为当前项目已入库记录；不能据此宣称全网覆盖。未判断记录单独计数。",
            ],
        ),
        ("执行摘要", finding_lines(narrative.executive_summary)),
        ("情感分布", [rows_table(["情感", "文章数", "占本批次比例"], sentiments)]),
    ]
    for dimension, title in [
        ("platform", "平台分布"),
        ("problem", "问题分类"),
        ("source", "来源账号分布"),
    ]:
        audit = query(audits, "breakdown", dimension)
        if audit:
            result.append(
                (
                    title,
                    [
                        rows_table(
                            ["类别", "文章数"],
                            [[g["label"], g["count"]] for g in audit["result"]["groups"]],
                        ),
                        f"查询 {audit['id']}；这是收录文章数量，不能视为账号影响力排名。",
                        f"展示 {len(audit['result']['groups'])} / "
                        f"{audit['result'].get('group_count', len(audit['result']['groups']))} 类；"
                        "超过展示上限的类别可通过数仓查询分页查看。",
                    ],
                )
            )
    return result + analysis_sections(narrative, audits, overview)


def analysis_sections(narrative, audits, overview):
    points = query(audits, "timeseries")["result"]["points"]
    return [
        (
            "入库热度与情感时间线",
            [
                rows_table(
                    ["时间（北京）", "情感", "文章数"],
                    [
                        [p["time"], LABELS.get(p["sentiment"], p["sentiment"]), p["count"]]
                        for p in points
                    ],
                ),
                "时间线依据入库时间。缺少此前连续区间与传播链，不能判断事件转折因果或预测后续热度。",
            ],
        ),
        ("观点与典型原文", finding_lines(narrative.viewpoints)),
        ("话题演变与趋势判断", finding_lines(narrative.evolution)),
        ("行动建议", finding_lines(narrative.recommendations)),
        (
            "局限与待补充数据",
            list(dict.fromkeys([*overview["limitations"], *narrative.limitations])),
        ),
    ]


def finding_lines(findings):
    return [
        f"{f.text}\n\n依据："
        + "、".join([*[f"文章 {a}" for a in f.article_ids], *[f"查询 {q}" for q in f.query_ids]])
        for f in findings
    ]


def evidence_sections(narrative, audits):
    cited = {
        a
        for items in (
            narrative.executive_summary,
            narrative.viewpoints,
            narrative.evolution,
            narrative.recommendations,
        )
        for finding in items
        for a in finding.article_ids
    }
    articles = {a["id"]: a for q in audits for a in q["result"].get("articles", [])}
    return [
        (
            f"原文证据 {key}",
            [
                articles[key]["title"],
                f"版本 {articles[key]['revision']} · {articles[key]['platform']} · "
                f"{articles[key]['source']}",
                articles[key]["quote"],
                "来源：" + safe_url(articles[key]["url"]),
                articles[key]["evidence_limit"],
            ],
        )
        for key in sorted(cited)
    ]


def safe_url(value):
    return value if urlparse(value).scheme in {"http", "https"} else "[无可访问链接]"


def render(narrative, dataset, audits):
    content = sections(narrative, dataset, audits) + evidence_sections(narrative, audits)
    charts = {"情感分布": chart(query(audits, "overview")["result"]["sentiments"], "情感分布 · 篇")}
    for dimension, title in [("platform", "平台分布"), ("problem", "问题分类")]:
        item = query(audits, "breakdown", dimension)
        if item:
            charts[title] = chart(item["result"]["groups"], title + " · 篇")
    timeline = defaultdict(int)
    for point in query(audits, "timeseries")["result"]["points"]:
        timeline[point["time"][:16].replace("T", " ")] += point["count"]
    charts["入库热度与情感时间线"] = chart(
        [{"label": k, "count": v} for k, v in timeline.items()], "入库时间分布 · 北京时间"
    )
    markdown = (
        "# "
        + narrative.title
        + "\n\n"
        + "\n\n".join("## " + title + "\n\n" + "\n\n".join(lines) for title, lines in content)
    )
    artifacts = [
        text_artifact("上汽通用舆情报告.md", "text/markdown", markdown),
        text_artifact(
            "上汽通用舆情报告.html", "text/html", html_report(narrative.title, content, charts)
        ),
        text_artifact(
            "数仓查询审计.json",
            "application/json",
            json.dumps(audits, ensure_ascii=False, indent=2),
        ),
    ]
    artifacts.extend(
        text_artifact(title + ".svg", "image/svg+xml", svg) for title, svg in charts.items()
    )
    artifacts.append(word_report(narrative.title, content))
    return artifacts


def text_artifact(name, media_type, content):
    return {"name": name, "mediaType": media_type, "content": content}


def html_report(title, content, charts):
    body = "".join(
        f"<section><h2>{html.escape(heading)}</h2>"
        + charts.get(heading, "")
        + "".join(html_block(line) for line in lines if line)
        + "</section>"
        for heading, lines in content
    )
    from ace_ext.reporting.template import PAGE_STYLE

    return (
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<meta http-equiv="Content-Security-Policy" '
        "content=\"default-src 'none'; style-src 'unsafe-inline'; img-src data:\">"
        f"<title>{html.escape(title)}</title><style>{PAGE_STYLE}</style>"
        "<main><p>ACE · PROJECT INTELLIGENCE / 项目舆情</p>"
        f"<h1>{html.escape(title)}</h1>{body}</main></html>"
    )


def html_block(line):
    if line.startswith("| "):
        rows = [row.strip("| ").split(" | ") for row in line.splitlines()]
        return (
            "<table>"
            + "".join(
                "<tr>" + "".join(f"<{tag}>{html.escape(cell)}</{tag}>" for cell in row) + "</tr>"
                for tag, row in [("th", rows[0]), *[("td", row) for row in rows[2:]]]
            )
            + "</table>"
        )
    if line.startswith("来源：http"):
        url = html.escape(line[3:], quote=True)
        return f'<p>来源：<a href="{url}" rel="noopener noreferrer">打开原文</a></p>'
    return "<p>" + html.escape(line) + "</p>"


def word_report(title, content):
    document = Document()
    document.styles["Normal"].font.name = "Microsoft YaHei"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(title, 0)
    for heading, lines in content:
        document.add_heading(heading, 1)
        for line in lines:
            if line.startswith("| "):
                rows = [row.strip("| ").split(" | ") for row in line.splitlines()]
                table = document.add_table(rows=0, cols=len(rows[0]))
                table.style = "Light Shading Accent 1"
                for row in [rows[0], *rows[2:]]:
                    for cell, value in zip(table.add_row().cells, row, strict=True):
                        cell.text = value
            elif line:
                document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return {
        "name": "上汽通用舆情报告.docx",
        "mediaType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "content_base64": base64.b64encode(stream.getvalue()).decode("ascii"),
    }
