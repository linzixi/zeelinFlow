from __future__ import annotations

import base64
from io import BytesIO
from pathlib import PurePath

from docx import Document
from docx.shared import Pt

from ace_ext.core.contracts import Envelope


def add_editable_documents(output: Envelope) -> Envelope:
    """将已验证的报告正文转为可编辑 Word，产物随节点快照持久保存。"""
    documents = []
    for artifact in output.artifacts:
        content = artifact.get("content")
        if not isinstance(content, str) or not content.strip():
            continue
        document = Document()
        document.styles["Normal"].font.name = "Microsoft YaHei"
        document.styles["Normal"].font.size = Pt(11)
        for line in content.splitlines():
            stripped = line.strip()
            if stripped.startswith("#"):
                level = min(len(stripped) - len(stripped.lstrip("#")), 3)
                document.add_heading(stripped.lstrip("# "), level=level)
            elif stripped.startswith(("- ", "* ")):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
        stream = BytesIO()
        document.save(stream)
        name = PurePath(str(artifact.get("name", "报告"))).stem + ".docx"
        documents.append(
            {
                "name": name,
                "kind": "report",
                "mediaType": "application/"
                "vnd.openxmlformats-officedocument.wordprocessingml.document",
                "content_base64": base64.b64encode(stream.getvalue()).decode("ascii"),
            }
        )
    output.artifacts.extend(documents)
    return output
